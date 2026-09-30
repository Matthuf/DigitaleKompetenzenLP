"""Wandelt die Item-Masterdatei (Excel) in items.json um. Einzige Datenquelle fuer HTML- und Backend-Version."""
import openpyxl, json, re, sys
from openpyxl.cell.rich_text import CellRichText, TextBlock
src = sys.argv[1] if len(sys.argv) > 1 else 'data/items_master.xlsx'
wb = openpyxl.load_workbook(src, rich_text=True)

def rich(v):
    """Zellinhalt als Text; fett formatierte Teile werden mit **…** markiert."""
    if v is None: return ''
    if not isinstance(v, CellRichText): return str(v).strip()
    out = ''
    for p in v:
        if isinstance(p, TextBlock):
            t = p.text
            if p.font is not None and p.font.b and t.strip():
                lead, trail = t[:len(t) - len(t.lstrip())], t[len(t.rstrip()):]
                out += lead + '**' + t.strip() + '**' + trail
            else: out += t
        else: out += str(p)
    return out.strip().replace('****', '')

def ss(t): return t.replace('ß', 'ss')
LEVELS = [("I","Einsteigen"),("II","Entdecken"),("III","Anwenden"),("IV","Integrieren"),("V","Entwickeln"),("VI","Weitergeben")]
DESC = {
 "1": "Digitale Medien für Kommunikation, Zusammenarbeit, Reflexion und die eigene Weiterbildung nutzen.",
 "2": "Digitale Materialien finden, erstellen, anpassen, ordnen und rechtlich korrekt teilen.",
 "3": "Digitale Medien im Unterricht einsetzen, Lernprozesse begleiten, kooperatives und selbstgesteuertes Lernen fördern.",
 "4": "Mit digitalen Mitteln Lernstände erheben, auswerten und Rückmeldungen geben.",
 "5": "Mit digitalen Medien Teilhabe ermöglichen, differenzieren und Lernende aktivieren.",
 "6": "Lernende befähigen, digitale Medien kompetent, kreativ und reflektiert zu nutzen.",
}
SHORT = {"1":"Berufsbezogenes Handeln","2":"Digitale Ressourcen","3":"Lehren und Lernen","4":"Lerndiagnose und Feedback","5":"Schülerorientierung","6":"Medienkompetenz der Lernenden"}
areas = []
for ws in wb.worksheets:
    m = re.match(r'^(\d)\. ', ws.title)
    if not m: continue
    aid = m.group(1)
    area = {"id": aid, "title": ws['A1'].value.split('. ',1)[1].strip(), "short": SHORT[aid], "description": DESC[aid], "subareas": []}
    cur = None
    # Spalte D: Originaltext (ALP Dillingen), Spalte E: Kurzfassung SZ (fett = das Konkrete der Stufe);
    # in der Kopfzeile eines Teilbereichs enthält Spalte E die Einleitung («Beispiele: …» / «Worum es geht: …»)
    for row in ws.iter_rows(min_row=4, values_only=True):
        a, b, c, d, e = (row + (None,)*5)[:5]
        if a and b is None and re.match(r'^\d\.\d ', str(a)):
            sid, title = str(a).split(' ', 1)
            cur = {"id": sid, "title": title.strip(), "ki": False, "intro": ss(rich(e).replace('**', '')), "levels": []}
            area["subareas"].append(cur)
        elif a and b and d:
            short = ss(rich(e))
            assert short, f"Kurzfassung fehlt bei {a}"
            assert short.count('**') % 2 == 0, f"Fettdruck unvollständig bei {a}"
            cur["levels"].append({"level": len(cur["levels"])+1, "roman": b, "label": c, "itemId": a,
                                  "text": short.replace('**', ''), "rich": short, "original": ss(str(d).strip())})
    for s in area["subareas"]:
        assert len(s["levels"]) == 6, s["id"]
    areas.append(area)
# Weiterbildung: Zuordnung Teilbereich -> fobizz-Themenbereiche (Blatt «Weiterbildung», optional)
PD_LEAD = "Schulentwicklung & Leadership"
themes = []
if 'Themen fobizz' in wb.sheetnames:
    themes = [r[0] for r in wb['Themen fobizz'].iter_rows(min_row=2, values_only=True) if r and r[0]]
if 'Weiterbildung' in wb.sheetnames:
    subs = {s["id"]: s for a in areas for s in a["subareas"]}
    for row in wb['Weiterbildung'].iter_rows(min_row=5, values_only=True):
        sid, _t, prim, sec = (row + (None,) * 4)[:4]
        if not sid or str(sid) not in subs: continue
        for th in (prim, sec):
            assert not th or not themes or th in themes, f"Unbekannter Themenbereich bei {sid}: {th}"
        subs[str(sid)]["pd"] = {"primary": prim or None, "secondary": sec or None}
out = {"instrument": "Digitale Kompetenzen Lehrpersonen Kanton Schwyz", "version": "SZ-2026.1-entwurf",
       "source": "DigCompEdu Bavaria, ALP Dillingen (Itemtexte ausgelesen am 28.09.2026); Kurzfassung Kanton Schwyz (Entwurf 30.09.2026)",
       "levels": [{"level": i+1, "roman": r, "label": l} for i, (r, l) in enumerate(LEVELS)],
       "areas": areas,
       "pd": {"provider": "fobizz", "lead": PD_LEAD, "themes": themes}}
json.dump(out, open('data/items.json', 'w'), ensure_ascii=False, indent=1)
print(len(areas), sum(len(a['subareas']) for a in areas))
